import csv
import pandas as pd
from openpyxl import load_workbook
import sqlite3

# ------------------Input ------------------

def input_data():
    data = []
    columns = input("Enter column names (comma-separated): ").strip().split(',')
    columns = [col.strip() for col in columns]

    print("Enter row data (type 'done' to finish):")
    while True:
        row_input = input(f"Enter values for {columns} (comma-separated): ")
        if row_input.lower() == 'done':
            break
        values = row_input.strip().split(',')
        if len(values) != len(columns):
            print("Number of values doesn't match columns.")
            continue
        data.append(dict(zip(columns, [v.strip() for v in values])))
    return data

# ------------------ CSV ------------------

def generate_csv(filename, data):
    if not data: 
        print("Data list is empty. Cannot generate CSV.")
        return
    with open(filename, mode='w', newline='') as file:
        writer = csv.DictWriter(file, fieldnames=data[0].keys())
        writer.writeheader()
        writer.writerows(data)

def csv_to_excel(csv_file, excel_file):
    df = pd.read_csv(csv_file)
    df.to_excel(excel_file, index=False)

# ------------------ SQL ------------------

def generate_sql(data, db_name, table_name):
    if not data:
        print("Data list is empty. Cannot generate SQL table.")
        return
    conn = sqlite3.connect(db_name)
    cursor = conn.cursor()
    
    # Create table with columns from data
    columns = data[0].keys()
    create_table_sql = f"CREATE TABLE IF NOT EXISTS {table_name} ({', '.join([f'{col} TEXT' for col in columns])})"
    cursor.execute(create_table_sql)
    
    # Insert data
    placeholders = ','.join(['?' for _ in columns])
    insert_sql = f"INSERT INTO {table_name} ({','.join(columns)}) VALUES ({placeholders})"
    for row in data:
        cursor.execute(insert_sql, list(row.values()))
    
    conn.commit()
    conn.close()
    print(f"SQL table '{table_name}' created in '{db_name}'")

def sql_to_csv(db_name, table_name, csv_file):
    conn = sqlite3.connect(db_name)
    df = pd.read_sql_query(f"SELECT * FROM {table_name}", conn)
    df.to_csv(csv_file, index=False)
    conn.close()
    print(f"Converted SQL table '{table_name}' to CSV: {csv_file}")

def sql_to_excel(db_name, table_name, excel_file):
    conn = sqlite3.connect(db_name)
    df = pd.read_sql_query(f"SELECT * FROM {table_name}", conn)
    df.to_excel(excel_file, index=False)
    conn.close()
    print(f"Converted SQL table '{table_name}' to Excel: {excel_file}")

class CSVEditor:
    def __init__(self, filename):
        self.filename = filename
        self.df = pd.read_csv(filename)

    def add_row(self):
        row = {col: input(f"Enter value for {col}: ") for col in self.df.columns}
        self.df.loc[len(self.df)] = row

    def delete_row(self):
        idx = int(input("Enter row index to delete: "))
        self.df = self.df.drop(idx).reset_index(drop=True)

    def add_column(self):
        col_name = input("Enter new column name: ")
        default_value = input("Enter default value: ")
        self.df[col_name] = default_value

    def update_cell(self):
        idx = int(input("Enter row index: "))
        col = input("Enter column name: ")
        val = input("Enter new value: ")
        self.df.at[idx, col] = val

    def save(self):
        self.df.to_csv(self.filename, index=False)
        print(f"Saved to {self.filename}")

    def show(self):
        print(self.df)

# ------------------ Excel------------------

def generate_excel(filename, data):
    df = pd.DataFrame(data)
    df.to_excel(filename, index=False)

def excel_to_csv(excel_file, csv_file):
    df = pd.read_excel(excel_file)
    df.to_csv(csv_file, index=False)

class ExcelEditor:
    def __init__(self, filename, sheet_name=None):
        self.filename = filename
        self.sheet_name = sheet_name or 'Sheet1'
        self.df = pd.read_excel(filename, sheet_name=self.sheet_name)

    def add_row(self):
        row = {col: input(f"Enter value for {col}: ") for col in self.df.columns}
        self.df.loc[len(self.df)] = row

    def delete_row(self):
        idx = int(input("Enter row index to delete: "))
        self.df = self.df.drop(idx).reset_index(drop=True)

    def add_column(self):
        col_name = input("Enter new column name: ")
        default_value = input("Enter default value: ")
        self.df[col_name] = default_value

    def update_cell(self):
        idx = int(input("Enter row index: "))
        col = input("Enter column name: ")
        val = input("Enter new value: ")
        self.df.at[idx, col] = val

    def save(self, new_sheet=None):
        sheet = new_sheet if new_sheet else self.sheet_name
        with pd.ExcelWriter(self.filename, engine='openpyxl', mode='a' if self.sheet_name else 'w') as writer:
            self.df.to_excel(writer, sheet_name=sheet, index=False)
        print(f"Saved to {self.filename} (sheet: {sheet})")

    def show(self):
        print(self.df)

# ------------------Menu ------------------

def main():
    while True:
        print("\n CSV, Excel & SQL Tool")
        print("1. Generate CSV")
        print("2. Generate Excel")
        print("3. Generate SQL")
        print("4. Convert CSV → Excel")
        print("5. Convert Excel → CSV")
        print("6. Convert SQL → CSV")
        print("7. Convert SQL → Excel")
        print("8. Edit CSV")
        print("9. Edit Excel")
        print("10. Exit")

        choice = input("Choose an option: ")

        if choice == '1':
            data = input_data()
            filename = input("Enter CSV filename (e.g. people.csv): ")
            generate_csv(filename, data)
            print("CSV file created.")

        elif choice == '2':
            data = input_data()
            filename = input("Enter Excel filename (e.g. people.xlsx): ")
            generate_excel(filename, data)
            print("Excel file created.")

        elif choice == '3':
            data = input_data()
            db_name = input("Enter database name (e.g. data.db): ")
            table_name = input("Enter table name: ")
            generate_sql(data, db_name, table_name)

        elif choice == '4':
            csv_file = input("Enter source CSV filename: ")
            excel_file = input("Enter target Excel filename: ")
            csv_to_excel(csv_file, excel_file)
            print("Converted CSV to Excel.")

        elif choice == '5':
            excel_file = input("Enter source Excel filename: ")
            csv_file = input("Enter target CSV filename: ")
            excel_to_csv(excel_file, csv_file)
            print("Converted Excel to CSV.")

        elif choice == '6':
            db_name = input("Enter database name (e.g. data.db): ")
            table_name = input("Enter table name: ")
            csv_file = input("Enter target CSV filename: ")
            sql_to_csv(db_name, table_name, csv_file)

        elif choice == '7':
            db_name = input("Enter database name (e.g. data.db): ")
            table_name = input("Enter table name: ")
            excel_file = input("Enter target Excel filename: ")
            sql_to_excel(db_name, table_name, excel_file)

        elif choice == '8':
            filename = input("Enter CSV filename to edit: ")
            editor = CSVEditor(filename)
            while True:
                print("\nCSV Editor")
                print("1. Show Data")
                print("2. Add Row")
                print("3. Delete Row")
                print("4. Add Column")
                print("5. Update Cell")
                print("6. Save and Exit")
                sub_choice = input("Select an option: ")
                if sub_choice == '1': editor.show()
                elif sub_choice == '2': editor.add_row()
                elif sub_choice == '3': editor.delete_row()
                elif sub_choice == '4': editor.add_column()
                elif sub_choice == '5': editor.update_cell()
                elif sub_choice == '6':
                    editor.save()
                    break

        elif choice == '9':
            filename = input("Enter Excel filename to edit: ")
            sheet = input("Enter sheet name (or leave blank for default): ") or None
            editor = ExcelEditor(filename, sheet)
            while True:
                print("\nExcel Editor")
                print("1. Show Data")
                print("2. Add Row")
                print("3. Delete Row")
                print("4. Add Column")
                print("5. Update Cell")
                print("6. Save and Exit")
                sub_choice = input("Select an option: ")
                if sub_choice == '1': editor.show()
                elif sub_choice == '2': editor.add_row()
                elif sub_choice == '3': editor.delete_row()
                elif sub_choice == '4': editor.add_column()
                elif sub_choice == '5': editor.update_cell()
                elif sub_choice == '6':
                    editor.save(sheet)
                    break

        elif choice == '10':
            print("Exiting tool. Goodbye!")
            break

        else:
            print("Invalid choice. Try again.")

if __name__ == "__main__":
    main()